import io
import os
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
INSTANCE_DIR.mkdir(exist_ok=True)
DATABASE_PATH = INSTANCE_DIR / "finance.db"

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", "dev-only-change-this-secret-key"),
    DATABASE=str(DATABASE_PATH),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    MAX_CONTENT_LENGTH=2 * 1024 * 1024,
)


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS roles (
            id_role INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            default_monthly_limit INTEGER NOT NULL DEFAULT 3000000
        );

        CREATE TABLE IF NOT EXISTS users (
            id_user INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role_id INTEGER NOT NULL,
            monthly_limit INTEGER NOT NULL DEFAULT 3000000,
            is_locked INTEGER NOT NULL DEFAULT 0 CHECK (is_locked IN (0, 1)),
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (role_id) REFERENCES roles(id_role)
        );

        CREATE TABLE IF NOT EXISTS transactions (
            id_transaction INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            transaction_date TEXT NOT NULL,
            transaction_type TEXT NOT NULL CHECK (transaction_type IN ('income', 'expense')),
            amount INTEGER NOT NULL CHECK (amount > 0),
            purpose TEXT NOT NULL,
            source TEXT NOT NULL CHECK (source IN ('personal', 'shared')),
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id_user) ON DELETE RESTRICT
        );

        CREATE INDEX IF NOT EXISTS idx_transactions_date ON transactions(transaction_date);
        CREATE INDEX IF NOT EXISTS idx_transactions_user_date ON transactions(user_id, transaction_date);
        CREATE INDEX IF NOT EXISTS idx_transactions_source_date ON transactions(source, transaction_date);
        """
    )
    db.execute(
        "INSERT OR IGNORE INTO roles (name, default_monthly_limit) VALUES (?, ?)",
        ("Admin", 0),
    )
    db.execute(
        "INSERT OR IGNORE INTO roles (name, default_monthly_limit) VALUES (?, ?)",
        ("User", 3000000),
    )
    admin_role = db.execute(
        "SELECT id_role FROM roles WHERE name = 'Admin'"
    ).fetchone()["id_role"]
    db.execute(
        """
        INSERT OR IGNORE INTO users
            (username, password_hash, full_name, role_id, monthly_limit, is_locked)
        VALUES (?, ?, ?, ?, ?, 0)
        """,
        (
            "admin",
            generate_password_hash("123"),
            "Quản trị viên",
            admin_role,
            0,
        ),
    )
    db.commit()


def get_csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


@app.before_request
def protect_post_requests():
    if request.method == "POST":
        expected = session.get("csrf_token", "")
        submitted = request.form.get("_csrf", "")
        if not expected or not secrets.compare_digest(expected, submitted):
            abort(400, description="Phiên biểu mẫu đã hết hạn. Hãy tải lại trang rồi thử lại.")


def current_user():
    if hasattr(g, "current_user"):
        return g.current_user

    user_id = session.get("user_id")
    if not user_id:
        g.current_user = None
        return None

    user = get_db().execute(
        """
        SELECT u.*, r.name AS role_name
        FROM users u
        JOIN roles r ON r.id_role = u.role_id
        WHERE u.id_user = ?
        """,
        (user_id,),
    ).fetchone()

    if user is None or user["is_locked"]:
        session.clear()
        flash("Tài khoản không còn hoạt động. Vui lòng đăng nhập lại hoặc liên hệ quản trị viên.", "warning")
        g.current_user = None
        return None

    g.current_user = user
    return user


def login_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if current_user() is None:
            return redirect(url_for("login"))
        return view(**kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped_view(**kwargs):
        if current_user()["role_name"] != "Admin":
            abort(403)
        return view(**kwargs)

    return wrapped_view


@app.context_processor
def inject_template_helpers():
    return {
        "current_user": current_user(),
        "csrf_token": get_csrf_token(),
        "today_iso": date.today().isoformat(),
    }


@app.template_filter("vnd")
def format_vnd(value):
    try:
        amount = float(value or 0)
    except (TypeError, ValueError):
        amount = 0
    return f"{amount:,.0f}".replace(",", ".") + " ₫"


@app.template_filter("date_vi")
def format_date_vi(value):
    if not value:
        return "—"
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return str(value)


def parse_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def parse_amount(value, minimum=1):
    try:
        amount = float(str(value).strip().replace(" ", "").replace(",", ""))
    except (TypeError, ValueError):
        return None
    if not amount.is_integer() or amount < minimum or amount > 1_000_000_000_000:
        return None
    return int(amount)


def get_filter_values(args):
    start = parse_date(args.get("start", ""))
    end = parse_date(args.get("end", ""))
    transaction_type = args.get("type", "").strip()
    source = args.get("source", "").strip()
    return {
        "start": start.isoformat() if start else "",
        "end": end.isoformat() if end else "",
        "type": transaction_type if transaction_type in ("income", "expense") else "",
        "source": source if source in ("personal", "shared") else "",
    }


def fetch_transactions(user, filters=None):
    filters = filters or {"start": "", "end": "", "type": "", "source": ""}
    db = get_db()
    sql = """
        SELECT t.*, u.username, u.full_name
        FROM transactions t
        JOIN users u ON u.id_user = t.user_id
        WHERE 1 = 1
    """
    params = []

    if user["role_name"] != "Admin":
        if filters.get("source") == "personal":
            sql += " AND t.source = 'personal' AND t.user_id = ?"
            params.append(user["id_user"])
        elif filters.get("source") == "shared":
            sql += " AND t.source = 'shared'"
        else:
            sql += " AND (t.source = 'shared' OR (t.source = 'personal' AND t.user_id = ?))"
            params.append(user["id_user"])

    if filters.get("start"):
        sql += " AND t.transaction_date >= ?"
        params.append(filters["start"])
    if filters.get("end"):
        sql += " AND t.transaction_date <= ?"
        params.append(filters["end"])
    if filters.get("type"):
        sql += " AND t.transaction_type = ?"
        params.append(filters["type"])
    if filters.get("source"):
        sql += " AND t.source = ?"
        params.append(filters["source"])

    sql += " ORDER BY t.transaction_date DESC, t.id_transaction DESC"
    return db.execute(sql, params).fetchall()


def sum_transactions(rows, transaction_type=None):
    return int(round(
        sum(
            int(row["amount"])
            for row in rows
            if transaction_type is None or row["transaction_type"] == transaction_type
        )
    ))


def wallet_balance(source, user_id=None):
    sql = """
        SELECT COALESCE(SUM(CASE WHEN transaction_type = 'income' THEN amount ELSE -amount END), 0)
        FROM transactions
        WHERE source = ?
    """
    params = [source]
    if user_id is not None:
        sql += " AND user_id = ?"
        params.append(user_id)
    value = get_db().execute(sql, params).fetchone()[0]
    return int(round(float(value or 0)))


def month_bounds(day=None):
    day = day or date.today()
    first = day.replace(day=1)
    if day.month == 12:
        next_month = date(day.year + 1, 1, 1)
    else:
        next_month = date(day.year, day.month + 1, 1)
    return first.isoformat(), (next_month - timedelta(days=1)).isoformat()


def month_spending(user_id, day=None):
    start, end = month_bounds(day)
    total = get_db().execute(
        """
        SELECT COALESCE(SUM(amount), 0)
        FROM transactions
        WHERE user_id = ? AND transaction_type = 'expense'
          AND transaction_date BETWEEN ? AND ?
        """,
        (user_id, start, end),
    ).fetchone()[0]
    return int(round(float(total or 0)))


def get_managed_transaction(transaction_id, user):
    row = get_db().execute(
        """
        SELECT t.*, u.username, u.full_name
        FROM transactions t JOIN users u ON u.id_user = t.user_id
        WHERE t.id_transaction = ?
        """,
        (transaction_id,),
    ).fetchone()
    if row is None:
        abort(404)
    if user["role_name"] != "Admin" and row["user_id"] != user["id_user"]:
        abort(403)
    return row


def validate_transaction_form():
    transaction_date = request.form.get("transaction_date", "").strip()
    transaction_type = request.form.get("transaction_type", "").strip()
    amount = parse_amount(request.form.get("amount", ""))
    purpose = request.form.get("purpose", "").strip()
    source = request.form.get("source", "").strip()

    parsed = parse_date(transaction_date)
    if not parsed:
        flash("Ngày giao dịch không hợp lệ.", "error")
    if transaction_type not in ("income", "expense"):
        flash("Vui lòng chọn loại giao dịch hợp lệ.", "error")
    if amount is None:
        flash("Số tiền phải là số nguyên VNĐ lớn hơn 0 và không vượt quá 1.000.000.000.000.", "error")
    if not purpose or len(purpose) > 180:
        flash("Mục đích giao dịch không được để trống và tối đa 180 ký tự.", "error")
    if source not in ("personal", "shared"):
        flash("Vui lòng chọn nguồn tiền hợp lệ.", "error")

    if not parsed or transaction_type not in ("income", "expense") or amount is None or not purpose or len(purpose) > 180 or source not in ("personal", "shared"):
        return None

    return {
        "transaction_date": parsed.isoformat(),
        "transaction_type": transaction_type,
        "amount": amount,
        "purpose": purpose,
        "source": source,
    }


@app.route("/")
@login_required
def dashboard():
    user = current_user()
    start, end = month_bounds()
    db = get_db()

    if user["role_name"] == "Admin":
        active_users = db.execute(
            """
            SELECT COUNT(*) FROM users u JOIN roles r ON r.id_role = u.role_id
            WHERE r.name = 'User' AND u.is_locked = 0
            """
        ).fetchone()[0]
        all_rows = fetch_transactions(user, {"start": start, "end": end, "type": "", "source": ""})
        all_time = fetch_transactions(user)
        summary = {
            "active_users": active_users,
            "monthly_expenses": sum_transactions(all_rows, "expense"),
            "shared_balance": wallet_balance("shared"),
            "transaction_count": len(all_time),
            "monthly_income": sum_transactions(all_rows, "income"),
        }
        recent = all_time[:8]
        return render_template("dashboard.html", summary=summary, recent=recent)

    my_month_rows = fetch_transactions(
        user,
        {"start": start, "end": end, "type": "", "source": "personal"},
    )
    own_spending = month_spending(user["id_user"])
    monthly_limit = int(user["monthly_limit"] or 0)
    limit_remaining = monthly_limit - own_spending
    shared_rows = fetch_transactions(
        user,
        {"start": start, "end": end, "type": "", "source": "shared"},
    )
    recent = fetch_transactions(user)[:8]
    summary = {
        "personal_balance": wallet_balance("personal", user["id_user"]),
        "shared_balance": wallet_balance("shared"),
        "monthly_spending": own_spending,
        "monthly_income": sum_transactions(my_month_rows, "income"),
        "monthly_limit": monthly_limit,
        "limit_remaining": limit_remaining,
        "limit_percent": min(100, round(own_spending / monthly_limit * 100)) if monthly_limit else 0,
        "shared_month_expenses": sum_transactions(shared_rows, "expense"),
    }
    return render_template("dashboard.html", summary=summary, recent=recent)


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user() is not None:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        user = get_db().execute(
            """
            SELECT u.*, r.name AS role_name
            FROM users u JOIN roles r ON r.id_role = u.role_id
            WHERE u.username = ?
            """,
            (username,),
        ).fetchone()

        if user and not user["is_locked"] and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id_user"]
            get_csrf_token()
            flash("Đăng nhập thành công. Chúc bạn quản lý tài chính hiệu quả!", "success")
            return redirect(url_for("dashboard"))

        flash("Tên đăng nhập hoặc mật khẩu không đúng, hoặc tài khoản đã bị khóa.", "error")

    return render_template("login.html")


@app.post("/logout")
@login_required
def logout():
    session.clear()
    flash("Bạn đã đăng xuất.", "success")
    return redirect(url_for("login"))


@app.route("/transactions")
@login_required
def transactions():
    user = current_user()
    filters = get_filter_values(request.args)
    rows = fetch_transactions(user, filters)
    totals = {
        "income": sum_transactions(rows, "income"),
        "expense": sum_transactions(rows, "expense"),
        "net": sum_transactions(rows, "income") - sum_transactions(rows, "expense"),
    }
    return render_template("transactions.html", rows=rows, filters=filters, totals=totals)


@app.route("/transactions/new", methods=["GET", "POST"])
@login_required
def new_transaction():
    if request.method == "POST":
        values = validate_transaction_form()
        if values:
            user = current_user()
            db = get_db()
            db.execute(
                """
                INSERT INTO transactions
                    (user_id, transaction_date, transaction_type, amount, purpose, source)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    user["id_user"],
                    values["transaction_date"],
                    values["transaction_type"],
                    values["amount"],
                    values["purpose"],
                    values["source"],
                ),
            )
            db.commit()
            if values["transaction_type"] == "expense" and user["role_name"] == "User":
                if month_spending(user["id_user"]) > int(user["monthly_limit"]):
                    flash("Đã lưu giao dịch, nhưng tổng chi trong tháng của tài khoản đã vượt hạn mức.", "warning")
                elif month_spending(user["id_user"]) >= int(user["monthly_limit"]) * 0.8:
                    flash("Đã lưu giao dịch. Lưu ý: bạn đã sử dụng ít nhất 80% hạn mức tháng.", "warning")
                else:
                    flash("Đã lưu giao dịch thành công.", "success")
            else:
                flash("Đã lưu giao dịch thành công.", "success")
            return redirect(url_for("transactions"))

    return render_template("transaction_form.html", transaction=None, page_title="Thêm giao dịch")


@app.route("/transactions/<int:transaction_id>/edit", methods=["GET", "POST"])
@login_required
def edit_transaction(transaction_id):
    user = current_user()
    transaction = get_managed_transaction(transaction_id, user)
    if request.method == "POST":
        values = validate_transaction_form()
        if values:
            get_db().execute(
                """
                UPDATE transactions
                SET transaction_date = ?, transaction_type = ?, amount = ?, purpose = ?, source = ?
                WHERE id_transaction = ?
                """,
                (
                    values["transaction_date"],
                    values["transaction_type"],
                    values["amount"],
                    values["purpose"],
                    values["source"],
                    transaction_id,
                ),
            )
            get_db().commit()
            flash("Đã cập nhật giao dịch.", "success")
            return redirect(url_for("transactions"))

    return render_template("transaction_form.html", transaction=transaction, page_title="Chỉnh sửa giao dịch")


@app.post("/transactions/<int:transaction_id>/delete")
@login_required
def delete_transaction(transaction_id):
    user = current_user()
    get_managed_transaction(transaction_id, user)
    get_db().execute("DELETE FROM transactions WHERE id_transaction = ?", (transaction_id,))
    get_db().commit()
    flash("Đã xóa giao dịch.", "success")
    return redirect(url_for("transactions"))


@app.route("/reports")
@login_required
def reports():
    user = current_user()
    period = request.args.get("period", "month")
    if period not in ("day", "week", "month", "year", "custom"):
        period = "month"
    reference = parse_date(request.args.get("reference_date", "")) or date.today()
    if period == "day":
        start, end = reference, reference
    elif period == "week":
        start = reference - timedelta(days=reference.weekday())
        end = start + timedelta(days=6)
    elif period == "year":
        start, end = date(reference.year, 1, 1), date(reference.year, 12, 31)
    elif period == "custom":
        start = parse_date(request.args.get("start", "")) or reference.replace(day=1)
        end = parse_date(request.args.get("end", "")) or reference
        if start > end:
            start, end = end, start
    else:
        start, end = month_bounds(reference)
        start, end = date.fromisoformat(start), date.fromisoformat(end)
        period = "month"

    filters = {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "type": "",
        "source": "",
    }
    rows = fetch_transactions(user, filters)
    income = sum_transactions(rows, "income")
    expense = sum_transactions(rows, "expense")
    purpose_totals = {}
    for row in rows:
        if row["transaction_type"] == "expense":
            purpose_totals[row["purpose"]] = purpose_totals.get(row["purpose"], 0) + int(row["amount"])
    purposes = sorted(purpose_totals.items(), key=lambda item: item[1], reverse=True)[:8]
    largest = max([amount for _, amount in purposes], default=0)
    purpose_chart = [
        {"purpose": purpose, "amount": amount, "percent": round(amount / largest * 100) if largest else 0}
        for purpose, amount in purposes
    ]
    return render_template(
        "reports.html",
        period=period,
        reference_date=reference.isoformat(),
        start=start.isoformat(),
        end=end.isoformat(),
        rows=rows,
        income=income,
        expense=expense,
        net=income - expense,
        purpose_chart=purpose_chart,
    )


def excel_safe_text(value):
    text_value = str(value if value is not None else "")
    if text_value.startswith(("=", "+", "-", "@")):
        return "'" + text_value
    return text_value


@app.route("/export.xlsx")
@login_required
def export_excel():
    user = current_user()
    filters = get_filter_values(request.args)
    rows = fetch_transactions(user, filters)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Giao dịch"
    headers = ["Ngày", "Loại giao dịch", "Thành viên", "Nguồn tiền", "Số tiền (VND)", "Mục đích"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B55")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(
            [
                row["transaction_date"],
                "Thu nhập" if row["transaction_type"] == "income" else "Chi tiêu",
                excel_safe_text(row["full_name"]),
                "Tiền chung" if row["source"] == "shared" else "Tiền cá nhân",
                int(row["amount"]),
                excel_safe_text(row["purpose"]),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [14, 18, 24, 18, 20, 42]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cells in sheet.iter_rows(min_row=2, min_col=5, max_col=5):
        cells[0].number_format = '#,##0'
    sheet.row_dimensions[1].height = 25

    summary = workbook.create_sheet("Tổng hợp")
    summary.append(["BÁO CÁO CHI TIÊU GIA ĐÌNH", ""])
    summary.append(["Thời gian bắt đầu", filters["start"] or "Tất cả"])
    summary.append(["Thời gian kết thúc", filters["end"] or "Tất cả"])
    summary.append(["Số giao dịch", len(rows)])
    summary.append(["Tổng thu (VND)", sum_transactions(rows, "income")])
    summary.append(["Tổng chi (VND)", sum_transactions(rows, "expense")])
    summary.append(["Chênh lệch thu - chi (VND)", sum_transactions(rows, "income") - sum_transactions(rows, "expense")])
    summary["A1"].font = Font(bold=True, size=14, color="1F6B55")
    for row_index in (5, 6, 7):
        summary.cell(row=row_index, column=2).number_format = '#,##0'
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 26

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = "bao_cao_chi_tieu_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def admin_users():
    db = get_db()
    user_role = db.execute("SELECT id_role FROM roles WHERE name = 'User'").fetchone()["id_role"]

    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        monthly_limit_value = parse_amount(request.form.get("monthly_limit", ""))
        valid_username = re.fullmatch(r"[a-z0-9_.-]{3,32}", username)

        if not valid_username:
            flash("Tên đăng nhập dài 3–32 ký tự, chỉ gồm chữ thường, số, dấu chấm, gạch dưới hoặc gạch ngang.", "error")
        elif not full_name or len(full_name) > 80:
            flash("Họ tên không được để trống và tối đa 80 ký tự.", "error")
        elif len(password) < 6:
            flash("Mật khẩu tài khoản thành viên phải có ít nhất 6 ký tự.", "error")
        elif monthly_limit_value is None:
            flash("Hạn mức tháng phải lớn hơn 0.", "error")
        else:
            try:
                db.execute(
                    """
                    INSERT INTO users (username, password_hash, full_name, role_id, monthly_limit, is_locked)
                    VALUES (?, ?, ?, ?, ?, 0)
                    """,
                    (
                        username,
                        generate_password_hash(password),
                        full_name,
                        user_role,
                        monthly_limit_value,
                    ),
                )
                db.commit()
                flash("Đã tạo tài khoản thành viên.", "success")
                return redirect(url_for("admin_users"))
            except sqlite3.IntegrityError:
                flash("Tên đăng nhập đã tồn tại.", "error")

    members = db.execute(
        """
        SELECT u.*, r.name AS role_name
        FROM users u JOIN roles r ON r.id_role = u.role_id
        WHERE r.name = 'User'
        ORDER BY u.created_at DESC, u.username COLLATE NOCASE
        """
    ).fetchall()
    start, end = month_bounds()
    spending_rows = db.execute(
        """
        SELECT user_id, COALESCE(SUM(amount), 0) AS total
        FROM transactions
        WHERE transaction_type = 'expense' AND transaction_date BETWEEN ? AND ?
        GROUP BY user_id
        """,
        (start, end),
    ).fetchall()
    spending_by_user = {row["user_id"]: float(row["total"]) for row in spending_rows}
    return render_template("admin_users.html", members=members, spending_by_user=spending_by_user)


@app.post("/admin/users/<int:user_id>/limit")
@admin_required
def update_user_limit(user_id):
    amount = parse_amount(request.form.get("monthly_limit", ""))
    member = get_db().execute(
        """
        SELECT u.id_user FROM users u JOIN roles r ON r.id_role = u.role_id
        WHERE u.id_user = ? AND r.name = 'User'
        """,
        (user_id,),
    ).fetchone()
    if member is None:
        abort(404)
    if amount is None:
        flash("Hạn mức phải lớn hơn 0.", "error")
    else:
        get_db().execute("UPDATE users SET monthly_limit = ? WHERE id_user = ?", (amount, user_id))
        get_db().commit()
        flash("Đã cập nhật hạn mức tháng.", "success")
    return redirect(url_for("admin_users"))


@app.post("/admin/users/<int:user_id>/toggle-lock")
@admin_required
def toggle_user_lock(user_id):
    db = get_db()
    member = db.execute(
        """
        SELECT u.id_user, u.is_locked FROM users u JOIN roles r ON r.id_role = u.role_id
        WHERE u.id_user = ? AND r.name = 'User'
        """,
        (user_id,),
    ).fetchone()
    if member is None:
        abort(404)
    new_status = 0 if member["is_locked"] else 1
    db.execute("UPDATE users SET is_locked = ? WHERE id_user = ?", (new_status, user_id))
    db.commit()
    flash("Đã mở khóa tài khoản." if new_status == 0 else "Đã khóa tài khoản.", "success")
    return redirect(url_for("admin_users"))


@app.route("/account/password", methods=["GET", "POST"])
@login_required
def change_password():
    user = current_user()
    if request.method == "POST":
        old_password = request.form.get("old_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        db_user = get_db().execute(
            "SELECT password_hash FROM users WHERE id_user = ?", (user["id_user"],)
        ).fetchone()

        if not check_password_hash(db_user["password_hash"], old_password):
            flash("Mật khẩu hiện tại không đúng.", "error")
        elif len(new_password) < 8:
            flash("Mật khẩu mới phải có ít nhất 8 ký tự.", "error")
        elif new_password != confirm_password:
            flash("Mật khẩu xác nhận không khớp.", "error")
        else:
            get_db().execute(
                "UPDATE users SET password_hash = ? WHERE id_user = ?",
                (generate_password_hash(new_password), user["id_user"]),
            )
            get_db().commit()
            flash("Đã đổi mật khẩu thành công.", "success")
            return redirect(url_for("dashboard"))

    return render_template("change_password.html")


@app.errorhandler(400)
def bad_request(error):
    return render_template("error.html", code=400, message=getattr(error, "description", "Yêu cầu không hợp lệ.")), 400


@app.errorhandler(403)
def forbidden(_error):
    return render_template("error.html", code=403, message="Bạn không có quyền truy cập chức năng này."), 403


@app.errorhandler(404)
def not_found(_error):
    return render_template("error.html", code=404, message="Không tìm thấy trang hoặc dữ liệu bạn yêu cầu."), 404


with app.app_context():
    init_db()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
